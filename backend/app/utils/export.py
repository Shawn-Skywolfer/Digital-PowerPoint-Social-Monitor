"""导出工具：文章/指标/评论/分析结果 → Excel / CSV。"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook

# 导出列定义：(字段key, 表头)
ARTICLE_COLUMNS = [
    ("title", "标题"), ("account_name", "公众号"), ("author", "作者"),
    ("publish_time", "发布时间"), ("read_num", "阅读量"), ("like_num", "点赞"),
    ("wow_num", "在看"), ("share_num", "转发"), ("collect_num", "收藏"),
    ("comment_count", "评论数"), ("status", "状态"), ("url", "链接"),
    ("content_text", "正文"),
]

COMMENT_COLUMNS = [
    ("article_title", "文章标题"), ("nickname", "昵称"), ("content", "评论内容"),
    ("like_num", "点赞"), ("is_sub", "是否二级"), ("comment_time", "评论时间"),
    ("status", "状态"),
]


# Excel 单元格上限 32767 字符，超出会被 openpyxl 拒绝
_CELL_MAX = 32767


def _fmt(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, str) and len(v) > _CELL_MAX:
        return v[:_CELL_MAX]
    return v if v is not None else ""


def to_csv(rows: list[dict], columns: list[tuple[str, str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([h for _, h in columns])
    for r in rows:
        writer.writerow([_fmt(r.get(k)) for k, _ in columns])
    # 加 BOM 便于 Excel 正确识别中文
    return buf.getvalue().encode("utf-8-sig")


def _write_sheet(ws, rows: list[dict], columns: list[tuple[str, str]]) -> None:
    ws.append([h for _, h in columns])
    for r in rows:
        ws.append([_fmt(r.get(k)) for k, _ in columns])
    # 简单自适应列宽
    for i, (_, h) in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            12, min(50, len(str(h)) * 4))


def to_xlsx(rows: list[dict], columns: list[tuple[str, str]],
            sheet_name: str = "数据") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    _write_sheet(ws, rows, columns)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_xlsx_multi(sheets: list[tuple[str, list[dict], list[tuple[str, str]]]]) -> bytes:
    """多 Sheet Excel：sheets = [(sheet名, 行数据, 列定义), ...]"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, columns in sheets:
        _write_sheet(wb.create_sheet(name), rows, columns)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_rows(rows: list[dict], columns: list[tuple[str, str]],
                fmt: str = "xlsx", sheet_name: str = "数据") -> tuple[bytes, str]:
    """返回 (bytes, 扩展名)。"""
    if fmt == "csv":
        return to_csv(rows, columns), "csv"
    return to_xlsx(rows, columns, sheet_name), "xlsx"
