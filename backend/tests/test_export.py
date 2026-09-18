"""导出接口测试：单篇分析按批次导出 + 聚合报告多 Sheet 导出。"""
import io

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models import (Account, AnalysisModule, AnalysisResult, Article,
                        CollectionReport, MetricSnapshot, ResultStatus)


def _seed():
    """造一套导出测试数据：账号+文章+指标+模块+两个批次的分析结果+聚合报告。"""
    db = SessionLocal()
    try:
        acc = Account(biz="export_test_biz", name="导出测试号")
        db.add(acc)
        db.flush()
        art1 = Article(account_id=acc.id, url="http://x/et1", title="导出测试文章一")
        art2 = Article(account_id=acc.id, url="http://x/et2", title="导出测试文章二")
        db.add_all([art1, art2])
        db.flush()
        db.add_all([MetricSnapshot(article_id=art1.id, read_num=100, like_num=10),
                    MetricSnapshot(article_id=art2.id, read_num=200, like_num=20)])
        mod = AnalysisModule(name="导出测试模块", target="content", prompt_template="t")
        db.add(mod)
        db.flush()
        db.add_all([
            AnalysisResult(article_id=art1.id, module_id=mod.id, model="m1",
                           result_text="结果A", tokens=11,
                           status=ResultStatus.success, batch_id="exp-batch-1"),
            AnalysisResult(article_id=art2.id, module_id=mod.id, model="m1",
                           result_text="结果B", tokens=22,
                           status=ResultStatus.success, batch_id="exp-batch-1"),
            AnalysisResult(article_id=art1.id, module_id=mod.id, model="m1",
                           result_text="结果C", tokens=33,
                           status=ResultStatus.success, batch_id="exp-batch-2"),
        ])
        report = CollectionReport(
            module_id=mod.id, model="m1", article_ids=[art1.id, art2.id],
            dimensions=[{"name": "关注焦点", "prompt": ""}],
            overview="本期聚焦储能话题。",
            result_text="raw", tokens=99, status=ResultStatus.success,
            result_json={
                "article_total": 2,
                "overview": "本期聚焦储能话题。",
                "topics": [{
                    "name": "储能安全", "summary": "多篇讨论储能电站安全",
                    "article_ids": [art1.id, art2.id], "article_count": 2,
                    "total_read": 300, "total_like": 30,
                    "top_articles": [
                        {"id": art2.id, "title": "导出测试文章二", "read_num": 200},
                        {"id": art1.id, "title": "导出测试文章一", "read_num": 100},
                    ],
                }],
                "dimension_insights": {"关注焦点": "储能与光伏"},
            },
        )
        db.add(report)
        pending = CollectionReport(
            module_id=mod.id, article_ids=[art1.id], status=ResultStatus.pending,
            batch_id="exp-pending")
        db.add(pending)
        db.commit()
        return report.id, pending.id
    finally:
        db.close()


def test_export_analysis_by_batch(client, auth_headers):
    _seed()
    # 全部导出：列里应含公众号
    resp = client.get("/api/v1/export/analysis", headers=auth_headers,
                      params={"fmt": "xlsx"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "公众号" in header and "分析结果" in header

    # 按批次导出：只含该批次的 2 行
    resp = client.get("/api/v1/export/analysis", headers=auth_headers,
                      params={"batch_id": "exp-batch-1", "fmt": "xlsx"})
    assert resp.status_code == 200
    assert "batch_exp-bat" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    texts = [row[header.index("分析结果")].value for row in ws.iter_rows(min_row=2)]
    assert sorted(texts) == ["结果A", "结果B"]


def test_export_collection(client, auth_headers):
    db = SessionLocal()
    try:
        report = (db.query(CollectionReport)
                  .filter(CollectionReport.status == ResultStatus.success).first())
        pending = (db.query(CollectionReport)
                   .filter(CollectionReport.batch_id == "exp-pending").first())
        rid, pid = report.id, pending.id
    finally:
        db.close()

    # 成功报告 → 三 Sheet 的 xlsx
    resp = client.get(f"/api/v1/export/collection/{rid}", headers=auth_headers)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["报告概览", "议题排行", "文章清单"]

    ws = wb["议题排行"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "储能安全" and rows[0][3] == 300  # 议题名 + 总阅读

    ov = {r[0]: r[1] for r in wb["报告概览"].iter_rows(min_row=2, values_only=True)}
    assert ov["总体概述"] == "本期聚焦储能话题。"
    assert ov["维度洞察·关注焦点"] == "储能与光伏"

    arts = list(wb["文章清单"].iter_rows(min_row=2, values_only=True))
    assert len(arts) == 2

    # CSV 退化为议题排行单表
    resp = client.get(f"/api/v1/export/collection/{rid}", headers=auth_headers,
                      params={"fmt": "csv"})
    assert resp.status_code == 200
    assert "储能安全" in resp.content.decode("utf-8-sig")

    # 未完成 → 400；不存在 → 404
    resp = client.get(f"/api/v1/export/collection/{pid}", headers=auth_headers)
    assert resp.status_code == 400
    resp = client.get("/api/v1/export/collection/999999", headers=auth_headers)
    assert resp.status_code == 404
